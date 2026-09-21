import json
import openpyxl
import os

def parse_val(val):
    if val is None:
        return None
    if isinstance(val, str):
        val = val.strip()
        if val == "":
            return None
    return val

def build_kulpes_json():
    # Source Excel file
    src_file = r"d:\Jay\vscode\attech0824\_archive\2_docs_and_notes\KULPES_产品目录_2026_V5.0 - 複製.xlsx"
    wb = openpyxl.load_workbook(src_file, data_only=True)
    
    products = []
    
    app_key_map = {
        "聚氨酯热熔胶": "pu_hot_melt_adhesive",
        "UV 树脂": "uv_resin",
        "聚氨酯弹性体": "pu_elastomer",
        "软包装": "flexible_packaging",
        "印刷油墨": "printing_inks",
        "粘合剂": "adhesives",
        "罐听涂料": "can_coatings",
        "热密封涂料": "heat_seal_coatings"
    }

    category_counts = {}

    for sheetname in wb.sheetnames:
        # Exclude test methods per user instruction
        if "测试方法" in sheetname:
            continue
        
        ws = wb[sheetname]
        category_name = sheetname.strip()
        category_counts[category_name] = 0
        
        # Subheaders in row 3 for application columns (from column 10 onwards)
        app_cols = []
        for c in range(10, ws.max_column + 1):
            h_val = ws.cell(3, c).value
            if h_val:
                h_clean = str(h_val).strip()
                app_cols.append((c, h_clean))
        
        for r in range(4, ws.max_row + 1):
            model = parse_val(ws.cell(r, 1).value)
            if not model:
                continue
            
            form = parse_val(ws.cell(r, 2).value)
            solid_content = parse_val(ws.cell(r, 3).value)
            tg = parse_val(ws.cell(r, 4).value)
            mw = parse_val(ws.cell(r, 5).value)
            ohv = parse_val(ws.cell(r, 6).value)
            av = parse_val(ws.cell(r, 7).value)
            mp = parse_val(ws.cell(r, 8).value)
            sp = parse_val(ws.cell(r, 9).value)
            
            # Application mapping
            applications = {}
            app_tags = []
            for col_idx, app_name in app_cols:
                app_flag = ws.cell(r, col_idx).value
                is_applicable = (app_flag == "+" or (isinstance(app_flag, str) and "+" in app_flag))
                en_key = app_key_map.get(app_name, app_name)
                applications[en_key] = is_applicable
                if is_applicable:
                    app_tags.append(app_name)
            
            # Form slug id
            slug_id = f"KULPES_{model.replace(' ', '_')}"
            
            prod_entry = {
                "id": slug_id,
                "brand": "KULPES®",
                "category": category_name,
                "model": model,
                "physical_form": form,
                "supply_form": {
                    "solid_content_percent": solid_content
                },
                "properties": {
                    "glass_transition_temperature_c": tg,
                    "molecular_weight_g_mol": mw,
                    "hydroxyl_value_mg_koh_g": str(ohv) if ohv is not None else None,
                    "acid_value_mg_koh_g": str(av) if av is not None else None,
                    "melting_point_c": mp,
                    "softening_point_c": sp
                },
                "applications": applications,
                "application_tags": app_tags
            }
            
            products.append(prod_entry)
            category_counts[category_name] += 1

    database = {
        "database_info": {
            "name": "KULPES® 树脂产品数据库 (排除测试方法)",
            "version": "2026 V5.0",
            "source_file": "KULPES_产品目录_2026_V5.0 - 複製.xlsx",
            "extracted_at": "2026-09-21",
            "company": {
                "name_zh": "昆勒聚合物(上海)有限公司",
                "name_en": "Kuller Polymer (Shanghai) Co., Ltd.",
                "address": "上海市闵行区都庄路2350号2幢501室",
                "postal_code": "201108",
                "phone": "021-6486 6300",
                "website": "http://www.kullerpolymer.com",
                "email": "kullerpolymer@163.com"
            },
            "excluded_sheets": ["测试方法"],
            "included_categories": list(category_counts.keys()),
            "category_counts": category_counts,
            "total_products": len(products)
        },
        "products": products
    }

    # Target paths for 37-product database (from copy file)
    target_path_1 = r"d:\Jay\vscode\attech0824\_archive\2_docs_and_notes\KULPES_产品数据库_2026_V5.0.json"
    target_path_2 = r"d:\Jay\vscode\attech0824\KULPES_产品数据库_2026_V5.0.json"
    
    for path in [target_path_1, target_path_2]:
        with open(path, "w", encoding="utf-8") as f:
            json.dump(database, f, ensure_ascii=False, indent=2)

    # Companion: also build full 70-product database from full catalog excel
    build_full_kulpes_json()
    print("JSON generation complete.")

def build_full_kulpes_json():
    src_file = r"d:\Jay\vscode\attech0824\KULPES_产品目录_2026_V5.0.xlsx"
    if not os.path.exists(src_file):
        return
    wb = openpyxl.load_workbook(src_file, data_only=True)
    
    # We can read from "全产品总表" which already normalized all 70 products!
    if "全产品总表" not in wb.sheetnames:
        return
    ws = wb["全产品总表"]
    
    full_products = []
    cat_counts = {}
    
    for r in range(3, ws.max_row + 1):
        idx = ws.cell(r, 1).value
        if not idx:
            continue
        cat = parse_val(ws.cell(r, 2).value)
        model = parse_val(ws.cell(r, 3).value)
        form = parse_val(ws.cell(r, 4).value)
        solid = parse_val(ws.cell(r, 5).value)
        solvent = parse_val(ws.cell(r, 6).value)
        tg = parse_val(ws.cell(r, 7).value)
        mw = parse_val(ws.cell(r, 8).value)
        ohv = parse_val(ws.cell(r, 9).value)
        av = parse_val(ws.cell(r, 10).value)
        mp = parse_val(ws.cell(r, 11).value)
        sp = parse_val(ws.cell(r, 12).value)
        visc = parse_val(ws.cell(r, 13).value)
        ph = parse_val(ws.cell(r, 14).value)
        apps_str = parse_val(ws.cell(r, 15).value) or ""
        
        app_list = [a.strip() for a in apps_str.split("、") if a.strip()]
        
        cat_counts[cat] = cat_counts.get(cat, 0) + 1
        
        full_products.append({
            "id": f"KULPES_{model.replace(' ', '_')}",
            "brand": "KULPES®",
            "category": cat,
            "model": model,
            "physical_form_or_monomer": form,
            "supply_form": {
                "solid_content_percent": solid,
                "solvent": solvent if solvent != "-" else None
            },
            "properties": {
                "glass_transition_temperature_c": tg,
                "molecular_weight_g_mol": mw,
                "hydroxyl_value_mg_koh_g": str(ohv) if ohv is not None else None,
                "acid_value_mg_koh_g": str(av) if av is not None else None,
                "melting_point_c": mp,
                "softening_point_c": sp,
                "viscosity_mpa_s": str(visc) if visc is not None else None,
                "ph_value": str(ph) if ph is not None and ph != "-" else None
            },
            "application_tags": app_list
        })
        
    full_db = {
        "database_info": {
            "name": "KULPES® 树脂产品全量数据库 (全6大系列，排除测试方法)",
            "version": "2026 V5.0",
            "source_file": "KULPES_产品目录_2026_V5.0.xlsx",
            "extracted_at": "2026-09-21",
            "company": {
                "name_zh": "昆勒聚合物(上海)有限公司",
                "name_en": "Kuller Polymer (Shanghai) Co., Ltd.",
                "address": "上海市闵行区都庄路2350号2幢501室",
                "postal_code": "201108",
                "phone": "021-6486 6300",
                "website": "http://www.kullerpolymer.com",
                "email": "kullerpolymer@163.com"
            },
            "excluded_sheets": ["测试方法与企业信息", "产品目录总览"],
            "total_categories": len(cat_counts),
            "category_counts": cat_counts,
            "total_products": len(full_products)
        },
        "products": full_products
    }
    
    with open(r"d:\Jay\vscode\attech0824\KULPES_产品数据库_全量70款_2026_V5.0.json", "w", encoding="utf-8") as f:
        json.dump(full_db, f, ensure_ascii=False, indent=2)
    with open(r"d:\Jay\vscode\attech0824\_archive\2_docs_and_notes\KULPES_产品数据库_全量70款_2026_V5.0.json", "w", encoding="utf-8") as f:
        json.dump(full_db, f, ensure_ascii=False, indent=2)

if __name__ == "__main__":
    build_kulpes_json()
